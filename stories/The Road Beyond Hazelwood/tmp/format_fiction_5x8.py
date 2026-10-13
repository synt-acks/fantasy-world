from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


source = Path('The Road Beyond Hazelwood.docx')
target = Path('The Road Beyond Hazelwood - 5x8 Fiction Format.docx')
doc = Document(source)

# Trade paperback trim size with a wider inner margin for binding.
for section in doc.sections:
    section.page_width = Inches(5)
    section.page_height = Inches(8)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.30)

# Mirror the asymmetric margins for facing pages in print.
settings = doc.settings._element
if settings.find(qn('w:mirrorMargins')) is None:
    settings.append(OxmlElement('w:mirrorMargins'))

normal = doc.styles['Normal']
normal.font.name = 'Garamond'
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Garamond')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Garamond')
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = Pt(13.2)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.widow_control = True

for para in doc.paragraphs:
    if para.style.name == 'Normal' and para.text.strip():
        fmt = para.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.line_spacing = Pt(13.2)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.widow_control = True
        if fmt.first_line_indent is not None and fmt.first_line_indent > 0:
            fmt.first_line_indent = Inches(0.22)
        for run in para.runs:
            run.font.name = 'Garamond'
            run.font.size = Pt(11)
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.rFonts
            if rfonts is None:
                rfonts = OxmlElement('w:rFonts')
                rpr.insert(0, rfonts)
            rfonts.set(qn('w:ascii'), 'Garamond')
            rfonts.set(qn('w:hAnsi'), 'Garamond')

# A restrained, centered title page that fits the narrower trim size.
title = doc.paragraphs[0]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.first_line_indent = Inches(0)
title.paragraph_format.space_before = Pt(164)
title.paragraph_format.space_after = Pt(0)
for run in title.runs:
    run.font.name = 'Garamond'
    run.font.size = Pt(24)
    run.font.bold = False
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Garamond')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Garamond')

# Chapter number and chapter title: centered, compact, and kept together.
for para in doc.paragraphs:
    if para.style.name == 'Heading 1':
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(44)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.keep_together = True
        for run in para.runs:
            run.font.name = 'Garamond'
            run.font.size = Pt(10.5)
            run.font.all_caps = True
            run.font.bold = False
            run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Garamond')
            run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Garamond')
    elif para.style.name == 'Heading 2':
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(25)
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.keep_together = True
        for run in para.runs:
            run.font.name = 'Garamond'
            run.font.size = Pt(14)
            run.font.italic = True
            run.font.bold = False
            run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Garamond')
            run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Garamond')

# Keep the page number understated and centered in the footer.
for section in doc.sections:
    for footer in (section.footer, section.first_page_footer):
        for para in footer.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.name = 'Garamond'
                run.font.size = Pt(9)

doc.save(target)
print(f'Wrote {target.resolve()} ({target.stat().st_size} bytes)')
