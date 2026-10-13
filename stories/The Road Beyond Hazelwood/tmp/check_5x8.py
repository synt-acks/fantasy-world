from pathlib import Path
import re

from docx import Document
from PIL import Image
from pypdf import PdfReader

root = Path('tmp/rendered_5x8')
pdf = next(root.glob('*.pdf'))
reader = PdfReader(str(pdf))
images = sorted(root.glob('page-*.png'), key=lambda p: int(re.search(r'(\d+)', p.stem).group(1)))
print('PDF pages:', len(reader.pages))
print('Page size in inches:', round(float(reader.pages[0].mediabox.width) / 72, 3), 'x', round(float(reader.pages[0].mediabox.height) / 72, 3))
print('Rendered PNGs:', len(images), 'dimensions:', Image.open(images[0]).size)
empty = [i + 1 for i, page in enumerate(reader.pages) if not (page.extract_text() or '').strip()]
print('Textless pages:', empty)
print('Title page text:', repr((reader.pages[0].extract_text() or '')[:120]))
print('Final page text:', repr((reader.pages[-1].extract_text() or '')[-160:]))
chapters = [(i + 1, (page.extract_text() or '').splitlines()[:4]) for i, page in enumerate(reader.pages) if 'CHAPTER ' in (page.extract_text() or '')]
print('Rendered chapter-opening pages:', chapters)
doc = Document('The Road Beyond Hazelwood - 5x8 Fiction Format.docx')
print('DOCX paragraphs:', len(doc.paragraphs))
print('Chapter heading paragraphs:', sum(p.style.name == 'Heading 1' for p in doc.paragraphs))
print('Section size/margins:', [(s.page_width.inches, s.page_height.inches, s.left_margin.inches, s.right_margin.inches, s.top_margin.inches, s.bottom_margin.inches) for s in doc.sections])
