from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent.parent
DOCX = ROOT / 'The Road Beyond Hazelwood - 4.25x6.87 Print Edition.docx'
PDF = ROOT / 'tmp' / 'rendered_425x687' / 'The Road Beyond Hazelwood - 4.25x6.87 Print Edition.pdf'

pdf = PdfReader(str(PDF))
page_texts = [page.extract_text() or '' for page in pdf.pages]
doc = Document(DOCX)
headings = [p for p in doc.paragraphs if p.style.name == 'Heading 1']
toc_entries = [p for p in doc.paragraphs if re.match(r'^\d+\.\s', p.text)]
if len(headings) != 28 or len(toc_entries) != 28:
    raise RuntimeError(f'Expected 28 chapters and TOC entries, got {len(headings)} and {len(toc_entries)}')

mapping = []
for heading, entry in zip(headings, toc_entries):
    needle = heading.text.upper()
    page_index = next((i for i, text in enumerate(page_texts) if needle in text), None)
    if page_index is None:
        raise RuntimeError(f'Could not find rendered chapter heading: {heading.text}')
    folios = re.findall(r'(?m)^\s*(\d{1,3})\s*$', page_texts[page_index])
    if not folios:
        raise RuntimeError(f'Could not find printed page number for {heading.text} on physical page {page_index + 1}')
    folio = folios[-1]
    entry.runs[-1].text = folio
    mapping.append((heading.text, page_index + 1, folio))

doc.save(DOCX)
print('Updated static contents page numbers:')
for item in mapping:
    print(*item, sep=' | ')
